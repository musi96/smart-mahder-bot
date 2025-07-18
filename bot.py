# Pillow compatibility workaround
import sys
if sys.version_info >= (3, 13):
    raise RuntimeError("Python 3.13 is not supported. Please use Python 3.10")

# Rest of your code remains the same...
import logging
from telegram import (
    Update,
    InlineKeyboardButton,
    InlineKeyboardMarkup
)
# ... rest of your existing bot.py code ...
from telegram.ext import (
    ApplicationBuilder,
    CommandHandler,
    CallbackQueryHandler,
    ContextTypes,
    MessageHandler,
    filters,
)
import os
import asyncio
from telegram.error import Conflict
import requests
import pdfplumber
from io import BytesIO
import python_pptx
from docx import Document

# --- Flask Web Server ---
from flask import Flask
import threading
app = Flask(__name__)
@app.route("/")
def health():
    return "OK", 200
def run_web():
    port = int(os.environ.get("PORT", 10000))
    app.run(host="0.0.0.0", port=port)
# -----------------------

# Configuration
BOT_TOKEN = "7969720988:AAHexLCWd8yMmQM7NiMyPhOmyCJ61fOXDwY"
CHANNEL_USERNAME = "HUESAchannel"
CHANNEL_ID = -1002040479523
ADMIN_ID = 1952017668
HF_TOKEN = "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"
TEXT_API = "https://api-inference.huggingface.co/models/gpt2"

# Menu Structure
MAIN_MENU = [
    ["📚 Educational Materials"],
    ["🧠 AI Features"],
    ["💬 Feedback"]
]

AI_FEATURES = [
    ["❓ Generate Questions"],
    ["📝 Create Summary"],
    ["🔙 Back"]
]

# All your course materials (complete structure)
courses = {
    "Economics": {
        "2 year": {
            "1 semester": [
                {
                    "name": "Calculus for Economics",
                    "files": [
                        {"file_id": "BQACAgQAAyEFAASeigO-AAMEaFh1xjxYUA9dOnrzT9gLMs7U_U8AAtEYAAKhpchSpSkYPNh2RDs2BA"},
                        {"file_id": "BQACAgQAAyEFAASeigO-AAMFaFh1xr6Nhb5I8lvjtmVVwfrrx1oAAtIYAAKhpchSu9ofq12HUWM2BA"}
                    ]
                },
                {
                    "name": "Accounting 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIBD2hZBmDMyIH9XkQhmXxeypNkYLkMAALgFgACAvDIUtsihO9VRhHeNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBEGhZBmDxZAPjdNnJ6T91cW6qs3qzAALhFgACAvDIUoM_a7ntjMW3NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBEWhZBmBFNDzyCgdcqXUkVEYrwhdBAALiFgACAvDIUtMv-EeSTSoPNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBEmhZBmDmiuOlJnIwF5oXC6seoAABHgAC4xYAAgLwyFK6NIdjWRBqgTYE"},
                        {"file_id": "BQACAgQAAxkBAAIBE2hZBmD0eV7kBJAfQDWeyIF7WrDBAALkFgACAvDIUn-VROTRFzCTNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBFGhZBmBxJATDi9CSGtc41IC_2MBWAALlFgACAvDIUr1Lc8-h9dUpNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBFWhZBmBxTaKe-jK-0g964Z2mCqoGAALmFgACAvDIUlOksdz9ddQdNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBFmhZBmBvjH6sYYrsGzILz0CNBGAEAALnFgACAvDIUizGebp3PbgNNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBF2hZBmCvvdQ2c_JSPGmNInO5zvn1AALoFgACAvDIUhX9FL6SQ9gyNgQ"}
                    ]
                },
                {
                    "name": "Macro Economics 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIByWhZn4PR-6a680joi0B3Aoz2bCbcAALcFwACAvDIUsqkj9_BXNpkNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBymhZn4PeIFppRYXYxUVEGYKaT_vGAALdFwACAvDIUsrlKC0k_7_HNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBy2hZn4Osn3iEPL1jwCiPZKna93ptAALeFwACAvDIUjOQsNfACtODNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBzGhZn4N_tqaFzqknoNl9v8qrA_HJAALfFwACAvDIUqJDGuXE5Nz1NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBzWhZn4PmmUKEpzM8ZJZ9zT9PgyexAALgFwACAvDIUqsW5LWBMWABNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBzmhZn4Pe_gSxz01JYyu8HGQeaOb5AALhFwACAvDIUrq-1e_01erdNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBz2hZn4N634WfAUJUpq4Y_--KKp2sAALiFwACAvDIUs228owIljvVNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB0GhZn4Nhz7wsoE2XAxiR68PbmgoaAALjFwACAvDIUi4AAc67BlBtOjYE"},
                        {"file_id": "BQACAgQAAxkBAAIB0WhZn4PlgtyYrY2G0ech4UCP5ZjDAALkFwACAvDIUvbZpImnYI4yNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB0mhZn4MuMEdd-anR-Wl5IFWAc4AOAALlFwACAvDIUmOZJ_Rlg7PZNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB3WhZn4rIWIwKvteA1S4OD-jkSz8vAALmFwACAvDIUqeLLt8SfX9qNgQ"}
                    ]
                },
                {
                    "name": "Micro Economics 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIB32hZoWvm6QhsxWnO8LC6aUVXNJ3nAALnFwACAvDIUutmFpX7hhoQNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB4GhZoWtrckPrkYfmJPrwTcSwot0FAALoFwACAvDIUm9ARP7GFhPfNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB4WhZoWv2PPbReP08Cb4WpDGMT5zpAALpFwACAvDIUiwqqXZjMmXQNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB4mhZoWtmgPYMCvg9TztJXtkVEHkKAALqFwACAvDIUrYpYhEaiAO9NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB42hZoWspeWj2rsL93KIIjCn05eaNAALrFwACAvDIUmFJzI88Bn5LNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB5GhZoWsge3pr2cvARP8wiAOQZPcaAALsFwACAvDIUthRNbMSa8kvNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB5WhZoWuXvimgHEEsMCm5dgQH2kcMAALtFwACAvDIUgbjsEi83uQTNgQ"}
                    ]
                },
                {
                    "name": "Introduction to Statistics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIB72hZo_uvtzHZAAI_swRRf9w4GNpcAALzFwACAvDIUjegO6N8lW5KNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB8GhZo_vIzcjLjG89xW_tS3RddSoJAAL0FwACAvDIUntqyRXdCF6MNgQ"}
                    ]
                },
                {
                    "name": "Basic Computer Skill of MS Applications",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIB_GhZqBF6uStb8p5Kp_7teinNgzkQAAL2FwACAvDIUhcHgVaUC6vcNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB_WhZqBEOLQq8Wywr4iXbfFWZcG27AAL3FwACAvDIUjww2Evy6-LrNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIB_mhZqBHMbgxzBgRLztXuOa9DXdMCAAL4FwACAvDIUtN127nEa3BdNgQ"}
                    ]
                }
            ],
            "2 semester": [
                {
                    "name": "Linear Algebra for Economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAICpWhZtCi-BC7w-UPDPa-wvJK_ObJRAAMYAAIC8MhSD1s-R5ltZc82BA"},
                        {"file_id": "BQACAgQAAxkBAAICpmhZtCgMZRWCs5DRR-wC9ZVLa7uiAAIBGAACAvDIUsGvKqPwj9x_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAICp2hZtCgDhrQthEnQBnGLNumDbftgAAIDGAACAvDIUthVIzX1Qqe3NgQ"},
                        {"file_id": "BQACAgQAAxkBAAICqGhZtCj1B43-4zKlUcmw8G7I6tCFAAICGAACAvDIUntoEf5o5LeENgQ"},
                        {"file_id": "BQACAgQAAxkBAAICqWhZtCgdEWW72I6iH1FHCfAgfb8HAAIEGAACAvDIUpkGeBqAPB2_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAICqmhZtCh8tkqq5Yaa7hpoNAABMzNhrAACBRgAAgLwyFIrL6EM7_3jCzYE"},
                        {"file_id": "BQACAgQAAxkBAAICq2hZtCjzNMXBKbE_HU4bbcNwTCOCAAIGGAACAvDIUvVrW4flTDk2NgQ"},
                        {"file_id": "BQACAgQAAxkBAAICrGhZtCi6LzwtGpEzq6LXjLzzlFVIAAIIGAACAvDIUoYpV99wd_KyNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICrWhZtCgfdR7e5_YDPRP6fPoY1JxCAAIJGAACAvDIUsNuRy0UmbHnNgQ"}
                    ]
                },
                {
                    "name": "Accounting 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAICt2hZtKp835jnDdaAp_UqkN20zxuXAAIKGAACAvDIUv2c5kRuqqNnNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICuGhZtKo_qHKS-rBPbYs-UpKh9ylCAAILGAACAvDIUh0Nzuq0mPwONgQ"},
                        {"file_id": "BQACAgQAAxkBAAICuWhZtKpujcb27-7GN-rsFwy1M-d-AAIMGAACAvDIUhU8Jjya_25rNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICumhZtKq9nDukHeMkZ6OyeH07OfF3AAINGAACAvDIUvZpUpMniBXGNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICu2hZtKr1Kd4of8q1xDRRbFHGA3eSAAIOGAACAvDIUpCutJF2ztGbNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICvGhZtKqPuaY5Vjo7XA3NF_Q3dUpjAAIPGAACAvDIUph7RdNcIV4SNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICvWhZtKoG2GYqOvmYB6BHd1w0XwbdAAIQGAACAvDIUj2VWfOIO9weNgQ"}
                    ]
                },
                {
                    "name": "Macro Economics 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAICxWhZtVoqzMQOEQ6hSI5463Zt4-YTAAIRGAACAvDIUs_rWIC4UrlhNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICxmhZtVpZvfszTcfQgNrQqZ4l0InwAAISGAACAvDIUjGvE1HgF-h6NgQ"},
                        {"file_id": "BQACAgQAAxkBAAICx2hZtVqjPqriD3LTxD4klxElJQSJAAITGAACAvDIUifxu8bQprJpNgQ"}
                    ]
                },
                {
                    "name": "Micro Economics 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAICy2hZtea6pN1CL5tWWmsUt-35gGhgAAIUGAACAvDIUi_lopYn4m7VNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICzGhZteaE9nwVEqBJHsgwFTBlPAiNAAIVGAACAvDIUqBJCxTzGAyPNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICzWhZtebZdw87_2Lo9nkGDCmoFCjvAAIWGAACAvDIUkgaSRbw6PLDNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICzmhZteZuFVPhQSa2YgofFo5bzeqoAAIXGAACAvDIUpzv8NUaoadmNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICz2hZteYRf78bXkZ6QuS4QkGAV4jlAAIYGAACAvDIUj8YDeqQXXsiNgQ"}
                    ]
                },
                {
                    "name": "Statistics for Economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIC1WhZtjY3UnlPzqJCzGTQ22hiqes3AALeGAACFaJBUfgkRvllkdxVNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIC1mhZtjYHP1Jt7Bhe-WdJUvrmlXolAALfGAACFaJBUT_W86L5NWyxNgQ"}
                    ]
                },
                {
                    "name": "Basic Writing Skill",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIC2WhZtqxpTOh-qhGKstxWro48RtUnAAIzCgACXBrYUmbTAuHahzXJNgQ"}
                    ]
                }
            ]
        },
        "3 year": {
            "1 semester": [
                {
                    "name": "Mathematical Economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDBWhZvM8eClDFYgo7CPDkJxQhP7kXAAJ5FgACDtQgU0w8DeCP0seINgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDBmhZvM95-qZflBrfHdJrS2q5sRyaAAL4DgAC1LGgUhlSv4Fttk4JNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDB2hZvM_HWjS8UpbAPyuQF3iotta5AAIOFgACVDwwU1lX7F-M15fENgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDCGhZvM8CWyqpoo5W80qrtvBIOwAB7wACDhgAAreaCVIadiY-n2rpODYE"},
                        {"file_id": "BQACAgQAAxkBAAIDCWhZvM_zaMTKPn5ipiT-MY4U01YWAAICFwAC91mhUfRvxlhsi09pNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDCmhZvM8feg3lFo6yYEAG8K_YbleFAAI9HQACw08YUUpQkVBlvMdENgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDC2hZvM9Rth5Kr01efUhEKRPWzoULAAIOGQACV-1oUOak6Nt2npK-NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDDGhZvM8XmOGkWYVJnrQCc20KjR3-AAKXHAAC_8IoUpHbVLM-Bu7hNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDDWhZvM_WEjweAAGngqPO4hCEURmLYwACPhcAAqkI6FJ0c83JNr2KkzYE"}
                    ]
                },
                {
                    "name": "Econometrics 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDF2hZvUInOerbVn7SXOKc7fuX_VrSAAJ1FgACDtQgU1cxMgTEGThoNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDGGhZvUILG5AZs9OheBAk-m9GbFPFAAKAFgACDtQgU-FlxpoRl1oUNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDGWhZvUKqMBOQ7hRQbOhAtnQ1vnBnAAMTAAIfJ1FT86d9Knfy2Bc2BA"},
                        {"file_id": "BQACAgQAAxkBAAIDGmhZvUJ309hYE82GXyLPQCexS4nvAAJrEwACmRqYUxrkuQUdY3apNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDG2hZvUJ2eXBk9mzOLi5aNwGcaK_QAAIHEwACjMXpUWyMm8hFPBekNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDHGhZvUKko638mBbQukgCgoYiSoTwAAJREwACjMXpUSs2c5B7Sd7fNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDHWhZvUKGKKpCDgHApN_ErOrzrdm3AAIKGAACt5oJUnyuqySQVSPUNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDHmhZvUJSKgVDUslymGRLpLGWMq7pAAL7FgAC91mhUXQ4q4a-7_vlNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDH2hZvUKPaauUqRhLAAGe1VUFhUPlFAACPB0AAsNPGFF5KQnANEgOJjYE"},
                        {"file_id": "BQACAgQAAxkBAAIDIGhZvUKBJvQFO3QIX1aLPn7S1rhZAAIoFwAC8HbwU-DYW5jyfm4-NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDIWhZvUJgGBf8-OSVORXLFsXH9v7bAAK2FgACAcc4UCbqEZtmVuPXNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDImhZvULzYto6ZqXwHOJijdOugg8QAAJfFQACpZbpUBDnzCv8evI7NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDI2hZvUJ3fwMh-RScgVIMNukro4CLAAJrLgACSyDQUT0xKqdNf7qDNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDJGhZvUIvHT3pxRtKE6dgZLqzgSBGAAKKFAACS9ygUBKxlMGiXm63NgQ"}
                    ]
                },
                {
                    "name": "Financial Economics 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDM2hZvaX3WtpOqL3xI40R8e8fy52-AAIgFQAClBwJU_B0VgKvZ-SdNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDNGhZvaX__B6Zi4fPf8D7ERX3tI3vAAI7FgACDdUAAVAjYK4iNEsGJzYE"},
                        {"file_id": "BQACAgQAAxkBAAIDNWhZvaXEs0VQ-V3v3uFxwvpdWE--AAIdEgACcGg5UTiN3rsgrIC_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDNmhZvaU67J_cxl6J-SQ3PF19AhS0AAIsEgACF4voURVxBD8b18v1NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDN2hZvaWFSYi0OdrcoaPpUTQhM_7dAAKeFQACSSKRUVN6BdNZlfnJNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDOGhZvaX4FNpjoBZIiYEQ4X99MXOkAAJ0FQACnighUmTBoOsK6Qz2NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDOWhZvaXvUszh0zUopOSpuCj3sTVUAAJ1FQACnighUonU6bcSmd2ENgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDOmhZvaWcdPE7rxnhbBHIR10DZddMAAJ2FQACnighUkk3Z-4Emo1QNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDO2hZvaWRvYxYtgTAHOP4kiQZuBFzAAJ4FQACnighUvLtr5oah9gNNgQ"}
                    ]
                },
                {
                    "name": "Introduction to Management",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDRWhZvgnFRmnycqh46TRggA6fDytqAAJ7EQACH8GAU153-r8JRyfBNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDRmhZvglxxFgJ6Gl2pQ5tg7RHpdEGAAKAEAACxWlwUKdlY3XWQUQ8NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDR2hZvgn33IiWzolDWXgRsnkLE8OaAAKBEAACxWlwUD5i4VG9o6ASNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDSGhZvgmpBZzAkem8HyiGHHPDJA7SAALpEQACfwABcFHjEMehOv7i8TYE"},
                        {"file_id": "BQACAgQAAxkBAAIDSWhZvgklTodRo1-5Mm5x7QysSbWnAAIlGAACcQ6xUGqhVdRgLPLENgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDSmhZvgkomD0RVFG2Jz8yUWHD8qsOAAL0GAACoOgZUqLA5mXpc5tnNgQ"}
                    ]
                },
                {
                    "name": "Labor Economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDUWhZvm6xLG9VnV9JF8kJy3U7IuGYAAJ6FgACDtQgU1kzFjXs6ce0NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDUmhZvm56JFe-vypIQQ5L2rMcIauUAAK5FAACtG8AAVOaheGQEkCP5zYE"},
                        {"file_id": "BQACAgQAAxkBAAIDU2hZvm7goqj5OKXLLvUAAcWcW77POQACuhQAArRvAAFTfZrGF1M9ZbM2BA"},
                        {"file_id": "BQACAgQAAxkBAAIDVGhZvm4BVZG-_Oefg9wjFWE9KjS0AAK7FAACtG8AAVPRSYwg50lSRzYE"},
                        {"file_id": "BQACAgQAAxkBAAIDVWhZvm7rpgjJlI8AAfwBOBxEqPhhegACPhMAAkoIkFPvwtisEhSrejYE"},
                        {"file_id": "BQACAgQAAxkBAAIDVmhZvm4AAUhI2rnfpiIVNxaAGTKIlgAC6RcAAtOGwVNCsTD5OArT1zYE"},
                        {"file_id": "BQACAgQAAxkBAAIDV2hZvm7DBEiVPpkqssCX0Sh5wLZHAALqFwAC04bBU_p2VLfjDjDjNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDWGhZvm5VYblel7BZOPg-dTq3_xI9AALsFwAC04bBU4M2eHdIzOLCNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDWWhZvm6e_6mGBsjI5_SKG2Pn2rM2AAJTGwACDgfhU88bEt7ga_16NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDWmhZvm73GDhLhRKmZ75ejrKRfjlmAAJUGwACDgfhUxAcbzRBf_VzNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDW2hZvm4524-Xk8MK06Pc790vrFoqAAJVGwACDgfhU-Fn76pyiAxRNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDXGhZvm6kDbrCtEaEopdLohuip3zYAAJWGwACDgfhU_AR63zRVyXVNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDXWhZvm5gKotHQjmmeMjUSy4TgP0IAAKGFgACipsBUAtU8HE5leWrNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDXmhZvm7OoatRMt6-181_Gf3MDQjMAAJnGAACPdphUXdNJRDL6emnNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDX2hZvm5g9Yd6oV5n-pfwQrgk2kOOAAJuGAACPdphUUKvpbzoywW6NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDYGhZvm5UuyktKIyMCo6opYY0dpKEAAK7FgACUXeoUSHZRETHB6FDNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDYWhZvm7XVnx7bS9uVTqGwF9j9bGnAALSFgACjse4UTh9TF0i4gXBNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDYmhZvm7PVf9AYmg1pd8NZGFaH8DbAAIMFgACjsfAUYcEL8ci1mRZNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDY2hZvm6eKHevOBqkcKR7066G_J3ZAAKrGAACiMTAUVoLediyfDcpNgQ"}
                    ]
                },
                {
                    "name": "Developmental Economics 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDd2hZvtdSln3q0KU7Ph7gMIrD9UVRAAIIEwACK3thUDjb8ByVuSyuNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDeGhZvtdhheTLeaKd8yLc0ygWS2LzAALLDQACMu_ZU7HZHaeihLddNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDeWhZvtdpLdl_Q0gmyjslUoUZ3DkjAAKBGAACLcqZU-G0orjJdxZ_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDemhZvtc6vgOUlmBbKtc2IxZSzJTlAAIQFgACVDwwUwpg0rHbzE4SNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDe2hZvtfolo3Mi5d1yHG7Vp2-q4vXAAIVFwAC91mhUXoWZHh8enspNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDfGhZvtelL7dGmmMRa1ATXnIPedfDAALWFAACmXB4UDgC38n9zCe0NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDfWhZvtclzUQx1iEmT1fcZ9SnOiVZAAJSEQACPcaxUvmUjEhrpIXwNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDfmhZvtd3zjY2q-5RT1IGKxdeIwToAAJbEAAC92j4UiPocX-E4ZgGNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDf2hZvtcLTq2q3ziWYHZv_y2EJKoJAAJ2FwACGTUhUGE92roWUGf_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDgGhZvtfpRqmTY6olSOOG11oviWvBAAJ3FwACGTUhUI9MCU5wjMNxNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDgWhZvtf4JAf3NAalDbSl1w1Lro7zAAJ4FwACGTUhUGF3baAZCAPmNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDgmhZvtcKfGOn7OJ-tSBQQhQ4YIobAAJQFQAC1HbIUeKAzFZZ-qr_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDg2hZvteluogiir6HJybDUsNwxUjAAAJRFQAC1HbIUbE_4HFDKM8MNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDhGhZvtda-b82JEzHMrJa5RSrGjwdAAJSFQAC1HbIUYZMZwThIOt8NgQ"}
                    ]
                },
                {
                    "name": "International Economics 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIDk2hZvz5DUmmLkiL1Jh6w0jNt8hGRAAJ7FgACDtQgU8bthk47hey4NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDlGhZvz6Djyjz2FMAAc49OW679_bcFAACfBYAAg7UIFPYyrPcg8rSBjYE"},
                        {"file_id": "BQACAgQAAxkBAAIDlWhZvz7ZKf7s1wFnmqAQ5DyqZaFgAAJ9FgACDtQgU44aPH4gX2vKNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDlmhZvz5jS_yM-oI1V7Z9-HDcbjtuAAJQEQACZN1pUkZD-1Xj9c80NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDl2hZvz6fJFqXDWX1VDYPi2Hpria0AAKuFAACOwPAUob7bgZIj7GuNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDmGhZvz6opelmI3VsK13AbbTkqbCkAAJMEAACVMGgU7UHsAxBZdxENgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDmWhZvz7zPyYIUNot0LT7yn3SkX5lAAILFAACrOk5URaV2iLP-swsNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDmmhZvz4bJPa54lYLAmUXySVL27MhAAIMFAACrOk5UT39SAQB4lSVNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDm2hZvz6AJLMxs6G3b35sXgdKfmKDAAINFAACrOk5UUfAl5UhhnWnNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDnGhZvz7vq4dWW5-cn2xWvkuzGFj4AALIFAACMLVhUxOtkjplMr6CNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDnWhZvz4gTA2cCP1U6hE-SgABMNGlYwACyRQAAjC1YVOgpEFjjs8rbzYE"},
                        {"file_id": "BQACAgQAAxkBAAIDnmhZvz5UoLGnMcI5qmNib6cRiHdIAALKFAACMLVhU3BcSxEKg9b9NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDn2hZvz7epivJsNId6RZqb136TMCAAALLFAACMLVhU13b8BGsM8gCNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDoGhZvz6tWQd6BhYSu5lGhzTTtSRCAALMFAACMLVhUyBfS67mXIvkNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDoWhZvz5XEvZGdToGhJBNxqV2Sy4TAAILGAACt5oJUgHqAtakdpn-NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDomhZvz5Ds442kwL-KD7bwmxh5HamAAIDFwAC91mhUeuN9_TWt5fhNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDo2hZvz4YRrf0luyuZvcbll6X6OXHAAI-HQACw08YUbCH9JkrsPHLNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDpGhZvz6kcqA4sDhyYE7ZMV07Qgo9AAI7HQACw08YUU3SFgU8jrZeNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDpWhZvz6JJWUCEuhPHOizEibE_QX7AAKUEgACXApYU1-CTEym5nErNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDpmhZvz4227zZ6GU_i-zzBAHaS7PqAAKVEgACXApYU2sAAQpG9exZAAE2BA"},
                        {"file_id": "BQACAgQAAxkBAAIDp2hZvz4Y4k3rJt3UwLHhddHJlvAxAAJPEQACZN1pUl9imsaX-f3iNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDqGhZvz7WBCjkIthsHYcyi0JkYZ8pAAIpFwACPk_wUy5S5NMvtF2UNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDqWhZvz5kiwFlE78z00Mq5BPozVyiAAIqFwACPk_wU893uNWNjOMMNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDqmhZvz58L1t5_aotS_PJ4p3LkRrdAAJAFwACbhBxUfOCYFBGjhvvNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIDq2hZvz6-FsZZtQmAKSWxj6_BmguPAALuGQAC7hJhUeBSw-YGhSefNgQ"}
                    ]
                }
            ],
            "2 semester": [
                {
                    "name": "Development economics 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEEmhZ0GOSJyU0XYCcmV8fNArVsRUVAAJcGAACJ9_5UuanxdZp6feHNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEE2hZ0GOMwhwvzl-drgqxFt_QZlXSAALaFwACky3BU5xTF8DpWESLNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEFGhZ0GNRsC3rWbQ4fcZvXMTJmGDrAALbFwACky3BU2QlWxFJaEO-NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEFWhZ0GMJ24-z-JMG2WhNJ-oW7-lqAALcFwACky3BU84v-FFtz1-iNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEFmhZ0GM8RuJonlaYjepMsEt0LORhAALdFwACky3BU51-ZIFE_4I3NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEF2hZ0GMdkexatNVOMsIyl9yx5RpKAALeFwACky3BU440OKj3vV1YNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEGGhZ0GNKl-lR40c5dsI2OsaCZ_GoAALfFwACky3BU8HxwETSFzpKNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEGWhZ0GNtxl-8SeR-qcT5LzTg5JSqAALhFwACky3BU-llX93ImwMXNgQ"}
                    ]
                },
                {
                    "name": "International economics 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEImhZ0P_85VBP3K0cYWXpuHU8-e9HAALvFQACczUoU2IiKQABH2PGqjYE"},
                        {"file_id": "BQACAgQAAxkBAAIEI2hZ0P-ohQKPqbMiEq_MP2bSWceOAALwFQACczUoU9NJZkzJ_jElNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEJGhZ0P-CZ4dvWjJK1FhJnY_knoEmAALxFQACczUoU--BTBKKBIPVNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEJWhZ0P85mxcS_L1qx2klAxHIwITkAALyFQACczUoU8pUrcNYNESENgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEJmhZ0P-TzRX3TsWCLxTQFB3j22LcAAJDGAAC4Hl4U2ilIxqNmc6YNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEJ2hZ0P_3NjMdJvO0HGe8-ZrtIog6AAJEGAAC4Hl4U6Tue9BAFNl8NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEKGhZ0P_k-SRJtRfgcLag7RkeO5KQAAIvFgACfq_JU-xjviIjR5NFNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEKmhZ0P9LIDvC1_o3HKj_ZpSI7gaVAAISGwACnOyxUKEpUUvYP19XNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEK2hZ0P9OqQ0HfiwQNAyPr0DwExl0AAIEFwAC91mhUWxtfCR5GupoNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIELGhZ0P9WezYlgwUy1vTr5k7oUG0pAAIdHAACC34hUbS4QDsvH5dCNgQ"}
                    ]
                },
                {
                    "name": "Econometrics 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEOGhZ0Vfk96NFgucJCLtA1v_qGlFCAAKHFwACt33ZU2XyWX-pMQu-NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEOWhZ0Vc9dGpe835LWRASWt9z6qwWAAL3FgACZjvZUyFsBZadPkGSNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEOmhZ0Vc_47oLRFMoOXY-MPc_4O4lAAKbFwACt33ZU37kdW8nP7SuNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEO2hZ0Vfqe9sH0DyJLQwCEEzCz06uAALpGAAC1T5wUoc1TopMk_yCNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEPGhZ0Vf40ikZ_jlDzI07pykR4NznAALqGAAC1T5wUo_STSvEwVLfNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEPmhZ0Vf8KhsH3fvdgNmFDUmTAlkDAAImGQACXsXpUG74TtoHUuIzNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEP2hZ0VfCfsswPt5UbXgswu9btj3aAAIWFgACanQ4UUqZibhXgGcNNgQ"}
                    ]
                },
                {
                    "name": "Research method for economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIESGhZ0bhzpk_6pBMWWCnhavLYK_-XAAJ4FAACNzogU3KOdDtzkldcNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIESWhZ0bg1Pous6gSFbi2sjRHIUXH9AAJ5FAACNzogUxgqWk7KriWDNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIESmhZ0bg9FD9mxaxFVqRA52im_7ZcAAJ6FAACNzogU8J8hjHa9he9NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIES2hZ0biDQJf3x9hXYqHem7hgPby1AAJ7FAACNzogU5JwWhWGC0HNNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIETGhZ0bjbZ22yB3E8ZJGqOQindmDlAAJ8FAACNzogU392d0k98TjVNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIETWhZ0bgQNoq66HKYfH0coRKzZs8YAAJ9FAACNzogU_HrvKt56DYyNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIETmhZ0bgXD5TUoEN65Mj7AR7fwKQUAAL4GAACZfaYUUUxhMfLi5WxNgQ"}
                    ]
                },
                {
                    "name": "Natural resources and environmental economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEVmhZ0h2OzIk-9x5OR7k1v6SiN1XcAALCFQACRDYpU_vG22trf22yNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEV2hZ0h3VOjrYxjnwuHfNbRe4rs1uAALDFQACRDYpU8eZq2l0JjoLNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEWGhZ0h0N8P1nsCkeegYcAVS6DSCqAALFFQACRDYpU_J3m5L5GXKLNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEWWhZ0h3bdAVH4kyW_nRoROmkgzcEAAKkFwACMn0xU9xRBIIQO3FUNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEWmhZ0h3xkJhzfwKBMVk0wpYOXdMGAAJHFwACsGRQUzXY14iMxYYqNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEW2hZ0h04tFiV8h6uuqmI60nWy4w1AAIbDwAC7WypUlegWqRv0Vw_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEXGhZ0h1207ztHtDfPJfozuZs4zdnAAKpGwACr_moU3xpOfPqZ-y_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEXWhZ0h2syRUPQbP2Ej46QAaBKBChAAJWHAAC41_oUCC4Q3OkBj9xNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEXmhZ0h1LAnmvtSc4hTvRQpEtpBEBAALIFgACSLwQUUdwiJI03_1ZNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEX2hZ0h1jWOq4AqWSTz0hJ86sZe3DAAIfGwACNT4wUc3uvHrA4Z-qNgQ"}
                    ]
                },
                {
                    "name": "Economics of industry",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEW2hZ0h04tFiV8h6uuqmI60nWy4w1AAIbDwAC7WypUlegWqRv0Vw_NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEa2hZ0nCGKcgznf8fyWjPTRo_q55cAAJzEwACbQVgUwfYj8cwYLlONgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEbGhZ0nAaIsJ03bMLUeBEh6ZoJgz_AAJOFwACksVBUYxahYGArijQNgQ"}
                    ]
                }
            ]
        },
        "4 year": {
            "1 semester": [
                {
                    "name": "Development planning and project analysis 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEpGhZ1_Wo2a-NjkQRJO8xMIPrLGXNAAJgGAACXNZRUnJNiMgdTibjNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEpWhZ1_UQuZRnvStKhesEeAInFt8iAAJhGAACXNZRUr8saqf1zVMkNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEpmhZ1_UZgjboDsdqUBgvGOqt2uj6AAJ3FwACGl6wUs5BqDUDGKbxNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEp2hZ1_VswR16AuRGYWT0TjBsEhP4AAIcGAACruQgU-xQ0-YtHHkWNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEqGhZ1_X84U4DCedYzzAc64uAdQOVAAIUEgAC6ev4UkPRlK78T-i7NgQ"}
                    ]
                },
                {
                    "name": "Statistical software application in economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEsGhZ2Em9TCaFHAIOhVZTQNAdm1u6AAIBFwAC4wzRUvVaHNBApDA5NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEsWhZ2EmdqvlYz9s2xl4p9ftXbvCMAAIGFwAC4wzRUpWV6sVI3b-LNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEsmhZ2El2AAHu_fWLYCCbxmDwDLpyRwACrBgAAg5u-FDwD1TvfoDB3zYE"}
                    ]
                },
                {
                    "name": "History of economic thought",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEtmhZ2IgoHzUtgE1YAuBtl9AP31m2AAK_EwACOtSAUlYeqsHvfYAnNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEt2hZ2Ii5yl166XjtklyVPiN1wQvrAALBEwACOtSAUiu_sZpQwJAzNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEuGhZ2Ii1tbx8l8xLXAVi2NHFi23yAALAEwACOtSAUtNIYHP-SxdjNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEuWhZ2IgC71LJ0mrf3R9o09UvaBi1AAMXAALjDNFS1BgAAUYY5egFNgQ"}
                    ]
                },
                {
                    "name": "Economics of agriculture",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEvmhZ2PUnxumUAAGsmzXN6WQ-tMxQegACKBUAAjv8sVIa5yjHPmGrWDYE"},
                        {"file_id": "BQACAgQAAxkBAAIEv2hZ2PXGV7SU2BJsph1FjDBsGsNRAAIIGAACLy7BUs_b6LT2uhShNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEwGhZ2PUZ5paKH-HyiVcGMZcqO_yfAAJuFgACgZ-BUAvbI3-HuKLaNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEwWhZ2PXUjUndR4Cv7sMj2NwHKUg4AAJvFgACgZ-BUJO8rUKWZ6SXNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEwmhZ2PWVPdbeWPzcwB5xdTxfM2JeAAL0GQACk1roUr8hyYmbxYbzNgQ"}
                    ]
                },
                {
                    "name": "Introduction to institutional and behavioral economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIEzGhZ2TrHrzIIGywCYuZyCZaXrEw0AAKfFAACbKXIUDtcT0eR5OQ2NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIEzWhZ2TqSKRxIPJH2Z5i5i0NE7Fq8AAKEGgAC0rEhUGwDjHYuerpVNgQ"}
                    ]
                },
                {
                    "name": "Monetary economics",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIE0GhZ2Xt5mh9hWnHXpc6wXfJYXD7QAAKREwACMGCAUuq94MUfuin9NgQ"}
                    ]
                }
            ],
            "2 semester": []
        }
    },
    "Accounting": {
        "2 year": {
            "1 semester": [
                {
                    "name": "Accounting 1",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAIBD2hZBmDMyIH9XkQhmXxeypNkYLkMAALgFgACAvDIUtsihO9VRhHeNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBEGhZBmDxZAPjdNnJ6T91cW6qs3qzAALhFgACAvDIUoM_a7ntjMW3NgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBEWhZBmBFNDzyCgdcqXUkVEYrwhdBAALiFgACAvDIUtMv-EeSTSoPNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBEmhZBmDmiuOlJnIwF5oXC6seoAABHgAC4xYAAgLwyFK6NIdjWRBqgTYE"},
                        {"file_id": "BQACAgQAAxkBAAIBE2hZBmD0eV7kBJAfQDWeyIF7WrDBAALkFgACAvDIUn-VROTRFzCTNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBFGhZBmBxJATDi9CSGtc41IC_2MBWAALlFgACAvDIUr1Lc8-h9dUpNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBFWhZBmBxTaKe-jK-0g964Z2mCqoGAALmFgACAvDIUlOksdz9ddQdNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBFmhZBmBvjH6sYYrsGzILz0CNBGAEAALnFgACAvDIUizGebp3PbgNNgQ"},
                        {"file_id": "BQACAgQAAxkBAAIBF2hZBmCvvdQ2c_JSPGmNInO5zvn1AALoFgACAvDIUhX9FL6SQ9gyNgQ"}
                    ]
                }
            ],
            "2 semester": [
                {
                    "name": "Accounting 2",
                    "files": [
                        {"file_id": "BQACAgQAAxkBAAICt2hZtKp835jnDdaAp_UqkN20zxuXAAIKGAACAvDIUv2c5kRuqqNnNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICuGhZtKo_qHKS-rBPbYs-UpKh9ylCAAILGAACAvDIUh0Nzuq0mPwONgQ"},
                        {"file_id": "BQACAgQAAxkBAAICuWhZtKpujcb27-7GN-rsFwy1M-d-AAIMGAACAvDIUhU8Jjya_25rNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICumhZtKq9nDukHeMkZ6OyeH07OfF3AAINGAACAvDIUvZpUpMniBXGNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICu2hZtKr1Kd4of8q1xDRRbFHGA3eSAAIOGAACAvDIUpCutJF2ztGbNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICvGhZtKqPuaY5Vjo7XA3NF_Q3dUpjAAIPGAACAvDIUph7RdNcIV4SNgQ"},
                        {"file_id": "BQACAgQAAxkBAAICvWhZtKoG2GYqOvmYB6BHd1w0XwbdAAIQGAACAvDIUj2VWfOIO9weNgQ"}
                    ]
                }
            ]
        }
    },
    "Gender": {},
    "Psychology": {},
    "Management": {},
    "PADM": {},
    "Sociology": {},
    "Journalism": {},
    "Hotel & Tourism Management": {}
}

logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s", 
    level=logging.INFO
)
logger = logging.getLogger(__name__)

async def is_user_member(user_id, context):
    try:
        member = await context.bot.get_chat_member(CHANNEL_ID, user_id)
        return member.status in ["member", "administrator", "creator"]
    except Exception as e:
        logger.warning(f"Failed to check membership: {e}")
        return False

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    user_id = update.effective_user.id
    if not await is_user_member(user_id, context):
        join_button = InlineKeyboardMarkup([
            [InlineKeyboardButton("Join Channel", url=f"https://t.me/{CHANNEL_USERNAME}")],
            [InlineKeyboardButton("✅ I've Joined", callback_data="check_membership")]
        ])
        await update.message.reply_text(
            "Please join our channel to access the bot!",
            reply_markup=join_button
        )
        return
    
    markup = InlineKeyboardMarkup([
        [InlineKeyboardButton(text[0], callback_data=text[0])] for text in MAIN_MENU
    ])
    await context.bot.send_message(
        chat_id=update.effective_chat.id,
        text="🏠 Main Menu:",
        reply_markup=markup
    )

async def show_fields(update: Update, context: ContextTypes.DEFAULT_TYPE):
    fields = [field for field in courses.keys() if courses[field]]  # Only show fields with content
    buttons = [[InlineKeyboardButton(field, callback_data=f"field|{field}")] for field in fields]
    markup = InlineKeyboardMarkup(buttons)
    await update.callback_query.edit_message_text(
        text="Select your field:",
        reply_markup=markup
    )

async def handle_course_selection(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    data = query.data.split('|')
    
    if data[0] == "field":
        field = data[1]
        years = list(courses[field].keys())
        buttons = [[InlineKeyboardButton(year, callback_data=f"year|{field}|{year}")] for year in years]
        markup = InlineKeyboardMarkup(buttons)
        await query.edit_message_text(
            text=f"Select year for {field}:",
            reply_markup=markup
        )
    
    elif data[0] == "year":
        field, year = data[1], data[2]
        semesters = list(courses[field][year].keys())
        buttons = [[InlineKeyboardButton(sem, callback_data=f"sem|{field}|{year}|{sem}")] for sem in semesters]
        markup = InlineKeyboardMarkup(buttons)
        await query.edit_message_text(
            text=f"Select semester for {field} - {year}:",
            reply_markup=markup
        )
    
    elif data[0] == "sem":
        field, year, sem = data[1], data[2], data[3]
        course_list = courses[field][year][sem]
        buttons = [[InlineKeyboardButton(course['name'], callback_data=f"course|{field}|{year}|{sem}|{idx}")] 
                  for idx, course in enumerate(course_list)]
        markup = InlineKeyboardMarkup(buttons)
        await query.edit_message_text(
            text=f"Select course for {field} - {year} - {sem}:",
            reply_markup=markup
        )
    
    elif data[0] == "course":
        field, year, sem, idx = data[1], data[2], data[3], int(data[4])
        course = courses[field][year][sem][idx]
        
        await context.bot.send_message(
            chat_id=query.message.chat_id,
            text=f"📚 {course['name']}"
        )
        
        for file in course['files']:
            await context.bot.send_document(
                chat_id=query.message.chat_id,
                document=file['file_id'],
                protect_content=True
            )
        
        # Show back button
        buttons = [[InlineKeyboardButton("🔙 Back", callback_data=f"sem|{field}|{year}|{sem}")]]
        await context.bot.send_message(
            chat_id=query.message.chat_id,
            text="What would you like to do next?",
            reply_markup=InlineKeyboardMarkup(buttons)
        )

async def button(update: Update, context: ContextTypes.DEFAULT_TYPE):
    query = update.callback_query
    await query.answer()
    user_id = query.from_user.id

    if not await is_user_member(user_id, context):
        join_button = InlineKeyboardMarkup([
            [InlineKeyboardButton("Join Channel", url=f"https://t.me/{CHANNEL_USERNAME}")],
            [InlineKeyboardButton("✅ I've Joined", callback_data="check_membership")]
        ])
        await query.edit_message_text(
            "Please join our channel to access the bot!",
            reply_markup=join_button
        )
        return

    if query.data == "check_membership":
        await start(query, context)
        return

    if query.data == "📚 Educational Materials":
        await show_fields(query, context)
    elif query.data == "🧠 AI Features":
        markup = InlineKeyboardMarkup([
            [InlineKeyboardButton(text[0], callback_data=text[0])] for text in AI_FEATURES
        ])
        await query.edit_message_text("🤖 AI Features:", reply_markup=markup)
    elif query.data == "🔙 Back":
        await start(query, context)
    elif query.data == "❓ Generate Questions":
        context.user_data['awaiting_file'] = True
        context.user_data['action'] = 'questions'
        await query.edit_message_text("📤 Please upload a file (PDF, PPT, DOCX, or TXT) to generate questions from")
    elif query.data == "📝 Create Summary":
        context.user_data['awaiting_file'] = True
        context.user_data['action'] = 'summary'
        await query.edit_message_text("📤 Please upload a file (PDF, PPT, DOCX, or TXT) to create a summary from")
    elif query.data.startswith(("field|", "year|", "sem|", "course|")):
        await handle_course_selection(update, context)

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if not context.user_data.get('awaiting_file'):
        return

    doc = update.message.document
    file = await context.bot.get_file(doc.file_id)
    file_bytes = BytesIO()
    await file.download_to_memory(file_bytes)
    file_bytes.seek(0)

    try:
        text = extract_text_from_file(file_bytes, doc.file_name)
        
        if context.user_data['action'] == 'questions':
            prompt = "Generate 10 quiz questions with answers about this content:"
        else:
            prompt = "Create a concise summary of this content:"
        
        response = generate_with_ai(text, prompt)
        
        # Split long messages
        for i in range(0, len(response), 4000):
            await update.message.reply_text(response[i:i+4000])
        
    except Exception as e:
        logger.error(f"Error processing file: {e}")
        await update.message.reply_text("❌ Sorry, I couldn't process that file. Please try another format.")
    
    finally:
        context.user_data.clear()
        await start(update, context)

def extract_text_from_file(file_bytes, filename):
    if filename.endswith('.pdf'):
        with pdfplumber.open(file_bytes) as pdf:
            return "\n".join([page.extract_text() for page in pdf.pages])
    elif filename.endswith('.docx'):
        doc = Document(file_bytes)
        return "\n".join([para.text for para in doc.paragraphs])
    elif filename.endswith('.pptx'):
        prs = python_pptx.Presentation(file_bytes)
        text = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text"):
                    text.append(shape.text)
        return "\n".join(text)
    elif filename.endswith('.txt'):
        return file_bytes.read().decode('utf-8')
    else:
        raise ValueError("Unsupported file format")

def generate_with_ai(text, prompt_prefix):
    headers = {"Authorization": f"Bearer {HF_TOKEN}"}
    payload = {
        "inputs": f"{prompt_prefix}\n\n{text}",
        "parameters": {"max_length": 1000}
    }
    response = requests.post(TEXT_API, headers=headers, json=payload)
    return response.json()[0]['generated_text']

async def feedback(update: Update, context: ContextTypes.DEFAULT_TYPE):
    context.user_data['awaiting_feedback'] = True
    await update.message.reply_text("💬 Please type your feedback message:")

async def handle_message(update: Update, context: ContextTypes.DEFAULT_TYPE):
    if context.user_data.get('awaiting_feedback'):
        feedback_text = update.message.text
        await context.bot.send_message(
            chat_id=ADMIN_ID,
            text=f"📩 Feedback from @{update.effective_user.username}:\n\n{feedback_text}"
        )
        await update.message.reply_text("✅ Thank you for your feedback!")
        context.user_data.clear()
    else:
        await start(update, context)

def main():
    try:
        app = ApplicationBuilder().token(BOT_TOKEN).build()
        app.add_handler(CommandHandler("start", start))
        app.add_handler(CallbackQueryHandler(button))
        app.add_handler(MessageHandler(filters.TEXT & ~filters.COMMAND, handle_message))
        app.add_handler(MessageHandler(filters.Document.ALL, handle_document))
        
        threading.Thread(target=run_web, daemon=True).start()
        app.run_polling()
    except Conflict as e:
        logger.error("Another instance is already running.")
    except Exception as e:
        logger.error(f"Error: {e}")

if __name__ == "__main__":
    main()
